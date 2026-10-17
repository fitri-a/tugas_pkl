import openpyxl
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

# Find the Excel file
data_dir = r'd:\pkl\progress_plscs\data'
files = os.listdir(data_dir)
print("Files in data dir:", files)

xlsx_file = None
for f in files:
    if f.endswith('.xlsx'):
        xlsx_file = os.path.join(data_dir, f)
        break

if not xlsx_file:
    print("No Excel file found!")
    sys.exit(1)

print(f"Opening: {xlsx_file}")
wb = openpyxl.load_workbook(xlsx_file, data_only=True)
sheets = wb.sheetnames
print('SHEETS:', sheets)
print()

def audit_sheet(ws, max_rows=20):
    print(f'Max Row: {ws.max_row}, Max Col: {ws.max_column}')
    for i, row in enumerate(ws.iter_rows(max_row=min(max_rows, ws.max_row), values_only=True)):
        print(f'  R{i+1}: {row}')

# Audit all sheets
for sheet_name in sheets:
    ws = wb[sheet_name]
    print(f'\n====== SHEET: {sheet_name} ======')
    audit_sheet(ws, 20)
