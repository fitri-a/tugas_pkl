"""
Service untuk impor dan ekspor file Excel
"""
import os
import logging
import re
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime, date
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

from app.database.connection import get_connection
from app.database.models import Progress
from app.services.progress_service import (
    add_progress, get_progress, get_rekap_mingguan,
    get_rekap_bulanan, hitung_status, format_angka, format_persen,
    get_settings
)
from app.config import MEJA_A, MEJA_B

logger = logging.getLogger(__name__)

# Header kolom DAY sheet
DAY_HEADERS = {
    'no': ['no', 'no.'],
    'tanggal': ['tanggal', 'date', 'tgl'],
    'nama_pic': ['nama pic', 'pic', 'nama', 'petugas'],
    'pekerjaan': ['pekerjaan', 'jenis pekerjaan', 'job'],
    'rincian': ['rincian pengerjaan', 'rincian', 'keterangan', 'detail'],
    'target': ['target'],
    'realisasi': ['realisasi', 'real', 'aktual'],
    'total_pct': ['total %', 'total persen', '%', 'persen', 'progress'],
    'status': ['keterangan', 'status', 'ket'],
}


# ============================================================
# IMPORT EXCEL
# ============================================================

def preview_excel_import(filepath: str) -> Dict[str, Any]:
    """
    Preview data sebelum diimpor.
    Return: {
        'sheets': list nama sheet yang valid,
        'rows_valid': jumlah baris valid,
        'rows_invalid': jumlah baris tidak valid,
        'errors': list pesan error,
        'preview_data': list dict baris preview,
        'duplicate_count': jumlah potensi duplikat,
    }
    """
    result = {
        'sheets': [],
        'rows_valid': 0,
        'rows_invalid': 0,
        'errors': [],
        'preview_data': [],
        'duplicate_count': 0,
        'filepath': filepath,
    }

    try:
        xl = pd.ExcelFile(filepath)
        available_sheets = xl.sheet_names
        logger.info(f"Import preview: sheets = {available_sheets}")
    except Exception as e:
        result['errors'].append(f"Gagal membuka file: {e}")
        return result

    target_sheets = {
        'DAY (MEJA A)': MEJA_A,
        'DAY (MEJA B)': MEJA_B,
    }

    seen = set()
    for sheet_name, meja in target_sheets.items():
        if sheet_name not in available_sheets:
            result['errors'].append(f"Sheet '{sheet_name}' tidak ditemukan.")
            continue

        result['sheets'].append(sheet_name)
        rows, errors, preview = _parse_day_sheet(xl, sheet_name, meja)
        result['rows_valid'] += len(rows)
        result['rows_invalid'] += len(errors)
        result['errors'].extend(errors)
        result['preview_data'].extend(preview[:50])
        for row in rows:
            key = (row.get('tanggal'), row.get('meja'), row.get('nama_pic'), row.get('pekerjaan'))
            if key in seen:
                result['duplicate_count'] += 1
            seen.add(key)

    return result


def import_excel_data(filepath: str, on_progress=None) -> Dict[str, Any]:
    """
    Mengimpor data dari file Excel ke database menggunakan transaksi.
    on_progress: callback(pct: int, pesan: str)
    """
    result = {
        'success': False,
        'rows_imported': 0,
        'rows_skipped': 0,
        'rows_failed': 0,
        'errors': [],
        'duplicates': [],
    }

    try:
        xl = pd.ExcelFile(filepath)
    except Exception as e:
        result['errors'].append(f"Gagal membuka file: {e}")
        return result

    target_sheets = {
        'DAY (MEJA A)': MEJA_A,
        'DAY (MEJA B)': MEJA_B,
    }

    all_rows = []
    for sheet_name, meja in target_sheets.items():
        if sheet_name not in xl.sheet_names:
            continue
        rows, errors, _ = _parse_day_sheet(xl, sheet_name, meja)
        all_rows.extend(rows)
        result['errors'].extend(errors)

    total = len(all_rows)
    conn = get_connection()

    try:
        conn.execute("BEGIN TRANSACTION")
        for i, row in enumerate(all_rows):
            try:
                # Cek duplikat
                from app.services.progress_service import check_duplicate
                is_dup = check_duplicate(
                    row['tanggal'], row['meja'], row['nama_pic'], row['pekerjaan']
                )
                if is_dup:
                    result['duplicates'].append(
                        f"{row['tanggal']} - {row['meja']} - {row['nama_pic']} - {row['pekerjaan']}"
                    )
                    result['rows_skipped'] += 1
                else:
                    p = Progress(**row)
                    conn.execute(
                        """INSERT INTO progress
                           (tanggal, meja, nama_pic, pekerjaan, rincian_pengerjaan,
                            target, realisasi, keterangan, imported_from)
                           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                        (p.tanggal, p.meja, p.nama_pic, p.pekerjaan,
                         p.rincian_pengerjaan, p.target, p.realisasi,
                         p.keterangan, os.path.basename(filepath))
                    )
                    result['rows_imported'] += 1

                if on_progress and total > 0:
                    on_progress(int((i + 1) / total * 100), f"Mengimpor baris {i+1} dari {total}...")

            except Exception as e:
                result['rows_failed'] += 1
                result['errors'].append(f"Baris {i+1}: {e}")

        conn.commit()
        result['success'] = True
    except Exception as e:
        conn.rollback()
        result['errors'].append(f"Transaksi gagal: {e}")
        result['success'] = False

    return result


def _parse_day_sheet(xl: pd.ExcelFile, sheet_name: str, meja: str) -> Tuple[List[Dict], List[str], List[Dict]]:
    """Mengurai sheet DAY dan mengembalikan (valid_rows, errors, preview)."""
    valid_rows = []
    errors = []
    preview = []

    try:
        df = xl.parse(sheet_name, header=None)
    except Exception as e:
        errors.append(f"[{sheet_name}] Gagal membaca: {e}")
        return valid_rows, errors, preview

    # Temukan baris header
    header_row_idx = _find_header_row(df)
    if header_row_idx is None:
        errors.append(f"[{sheet_name}] Baris header tidak ditemukan.")
        return valid_rows, errors, preview

    # Set header
    df.columns = [str(c).strip().lower() if c else '' for c in df.iloc[header_row_idx]]
    df = df.iloc[header_row_idx + 1:].reset_index(drop=True)

    # Temukan kolom
    col_map = _map_columns(df.columns.tolist())
    if not col_map.get('tanggal') or not col_map.get('target') or not col_map.get('realisasi'):
        errors.append(f"[{sheet_name}] Kolom penting tidak ditemukan: {df.columns.tolist()}")
        return valid_rows, errors, preview

    for idx, row in df.iterrows():
        try:
            tanggal_raw = row.get(col_map.get('tanggal'), None)
            if tanggal_raw is None or (isinstance(tanggal_raw, float) and pd.isna(tanggal_raw)):
                continue

            tanggal_str = _parse_date(tanggal_raw)
            if not tanggal_str:
                errors.append(f"[{sheet_name}] Baris {idx+2}: Format tanggal tidak valid: {tanggal_raw}")
                continue

            nama_pic = str(row.get(col_map.get('nama_pic', ''), '') or '').strip()
            pekerjaan = str(row.get(col_map.get('pekerjaan', ''), '') or '').strip()
            rincian = str(row.get(col_map.get('rincian', ''), '') or '').strip()
            target_raw = row.get(col_map.get('target', ''), None)
            realisasi_raw = row.get(col_map.get('realisasi', ''), None)

            target = _parse_number(target_raw)
            realisasi = _parse_number(realisasi_raw)

            if not pekerjaan:
                continue  # Lewati baris kosong

            data = {
                'tanggal': tanggal_str,
                'meja': meja,
                'nama_pic': nama_pic or 'Tidak Diketahui',
                'pekerjaan': pekerjaan,
                'rincian_pengerjaan': rincian,
                'target': target,
                'realisasi': realisasi,
                'keterangan': '',
                'imported_from': '',
            }
            valid_rows.append(data)
            if len(preview) < 50:
                preview.append({**data, '_sheet': sheet_name})

        except Exception as e:
            errors.append(f"[{sheet_name}] Baris {idx+2}: {e}")

    return valid_rows, errors, preview


def _find_header_row(df: pd.DataFrame) -> Optional[int]:
    """Menemukan baris header berdasarkan kata kunci."""
    keywords = ['tanggal', 'target', 'realisasi', 'pic', 'pekerjaan', 'date']
    for i, row in df.iterrows():
        row_vals = [str(v).lower().strip() for v in row if v and not pd.isna(v) if isinstance(v, str)]
        matches = sum(1 for k in keywords if any(k in v for v in row_vals))
        if matches >= 2:
            return int(i)
    return None


def _map_columns(columns: List[str]) -> Dict[str, str]:
    """Memetakan nama kolom ke field standar."""
    result = {}
    for col in columns:
        col_lower = col.lower().strip()
        for field, keywords in DAY_HEADERS.items():
            if any(kw in col_lower for kw in keywords) and field not in result:
                result[field] = col
                break
    return result


def _parse_date(val) -> Optional[str]:
    """Mengurai nilai tanggal dari Excel ke ISO string."""
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime('%Y-%m-%d')
    if isinstance(val, date):
        return val.strftime('%Y-%m-%d')
    val_str = str(val).strip()
    # Coba berbagai format
    formats = ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%d %m %Y', '%Y/%m/%d']
    for fmt in formats:
        try:
            return datetime.strptime(val_str, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    # Coba parse dari Excel serial number
    try:
        n = float(val_str)
        if 1 < n < 100000:
            from datetime import timedelta
            base = datetime(1899, 12, 30)
            d = base + timedelta(days=n)
            return d.strftime('%Y-%m-%d')
    except (ValueError, TypeError):
        pass
    return None


def _parse_number(val) -> Optional[float]:
    """Mengurai nilai angka dari Excel."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        if pd.isna(val):
            return None
        return float(val)
    val_str = str(val).strip().replace(' ', '')
    # Excel numerik biasanya sudah berupa angka; dukung juga format lokal 1.250,50.
    if ',' in val_str and '.' in val_str:
        if val_str.rfind(',') > val_str.rfind('.'):
            val_str = val_str.replace('.', '').replace(',', '.')
        else:
            val_str = val_str.replace(',', '')
    elif ',' in val_str:
        val_str = val_str.replace(',', '.')
    try:
        return float(val_str)
    except ValueError:
        return None


# ============================================================
# EXPORT EXCEL
# ============================================================

def export_to_excel(
    filepath: str,
    meja: str = None,
    tanggal_mulai: str = None,
    tanggal_selesai: str = None,
    nama_pic: str = None,
    include_weekly: bool = True,
    include_monthly: bool = True,
) -> bool:
    """Ekspor data ke file Excel dengan format profesional."""
    try:
        wb = Workbook()
        settings = get_settings()
        threshold_pct = settings.achievement_threshold * 100

        # Sheet 1: Data Harian
        ws_daily = wb.active
        ws_daily.title = "Data Harian"
        data = get_progress(
            meja=meja,
            tanggal_mulai=tanggal_mulai,
            tanggal_selesai=tanggal_selesai,
            nama_pic=nama_pic,
            order_by='tanggal ASC, meja ASC'
        )
        _write_daily_sheet(ws_daily, data, threshold_pct)

        # Sheet 2: Rekap Mingguan (jika diperlukan)
        if include_weekly and tanggal_mulai and tanggal_selesai:
            ws_week = wb.create_sheet("Rekap Mingguan")
            rekap_week = get_rekap_mingguan(tanggal_mulai, tanggal_selesai, meja)
            _write_weekly_sheet(ws_week, rekap_week, tanggal_mulai, tanggal_selesai, threshold_pct)

        # Sheet 3: Rekap Bulanan
        if include_monthly:
            ws_month = wb.create_sheet("Rekap Bulanan")
            _write_monthly_sheet_all(ws_month, meja, settings.report_year, threshold_pct)

        # Sheet 4: Statistik
        ws_stats = wb.create_sheet("Statistik")
        _write_stats_sheet(ws_stats, settings)

        wb.save(filepath)
        logger.info(f"Export berhasil: {filepath}")
        return True

    except Exception as e:
        logger.error(f"Export gagal: {e}")
        raise


def _write_daily_sheet(ws, data: List[Progress], threshold_pct: float):
    """Menulis sheet data harian."""
    # Judul
    ws['A1'] = 'LAPORAN PROGRESS PLSCS HARIAN'
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1A3A5C', end_color='1A3A5C', fill_type='solid')
    ws.merge_cells('A1:J1')
    ws['A1'].alignment = Alignment(horizontal='center')

    # Header
    headers = ['No', 'Tanggal', 'Meja', 'Nama PIC', 'Pekerjaan',
               'Rincian Pengerjaan', 'Target', 'Realisasi', 'Progress %', 'Status']
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')

    # Data
    STATUS_COLORS = {
        'TERCAPAI': '22C55E',
        'TIDAK TERCAPAI': 'EF4444',
        'BELUM DIISI': 'F59E0B',
        'DATA TIDAK LENGKAP': '94A3B8',
    }

    for i, p in enumerate(data, 1):
        row = i + 2
        pct = p.progress_pct
        status = hitung_status(p.realisasi, p.target, threshold_pct)
        color = STATUS_COLORS.get(status, 'FFFFFF')

        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=_fmt_date(p.tanggal))
        ws.cell(row=row, column=3, value=p.meja)
        ws.cell(row=row, column=4, value=p.nama_pic)
        ws.cell(row=row, column=5, value=p.pekerjaan)
        ws.cell(row=row, column=6, value=p.rincian_pengerjaan)
        ws.cell(row=row, column=7, value=p.target)
        ws.cell(row=row, column=8, value=p.realisasi)
        pct_cell = ws.cell(row=row, column=9, value=round(pct / 100, 4) if pct else None)
        pct_cell.number_format = '0.00%'
        status_cell = ws.cell(row=row, column=10, value=status)
        status_cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        status_cell.font = Font(color='FFFFFF', bold=True)
        status_cell.alignment = Alignment(horizontal='center')

        # Zebra stripe
        if i % 2 == 0:
            for col in range(1, 10):
                c = ws.cell(row=row, column=col)
                if not c.fill or c.fill.fgColor.rgb == '00000000':
                    c.fill = PatternFill(start_color='F0F4F8', end_color='F0F4F8', fill_type='solid')

    # Lebar kolom
    col_widths = [5, 12, 10, 25, 30, 40, 12, 12, 12, 16]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _write_weekly_sheet(ws, data: List[Dict], tgl_mulai: str, tgl_selesai: str, threshold_pct: float):
    """Menulis sheet rekap mingguan."""
    ws['A1'] = f'REKAP MINGGUAN {tgl_mulai} s/d {tgl_selesai}'
    ws['A1'].font = Font(bold=True, size=12, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1A3A5C', end_color='1A3A5C', fill_type='solid')
    ws.merge_cells('A1:H1')

    headers = ['No', 'Meja', 'PIC', 'Pekerjaan', 'Total Target', 'Total Realisasi', 'Progress %', 'Status']
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center')

    for i, d in enumerate(data, 1):
        row = i + 2
        pct = (d['total_realisasi'] / d['total_target'] * 100) if d['total_target'] > 0 else 0
        status = hitung_status(d['total_realisasi'], d['total_target'], threshold_pct)
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=d.get('meja', ''))
        ws.cell(row=row, column=3, value=d.get('nama_pic', ''))
        ws.cell(row=row, column=4, value=d.get('pekerjaan', ''))
        ws.cell(row=row, column=5, value=d.get('total_target', 0))
        ws.cell(row=row, column=6, value=d.get('total_realisasi', 0))
        c = ws.cell(row=row, column=7, value=round(pct / 100, 4))
        c.number_format = '0.00%'
        ws.cell(row=row, column=8, value=status)


def _write_monthly_sheet_all(ws, meja: str, tahun: int, threshold_pct: float):
    """Menulis sheet rekap bulanan semua bulan."""
    ws['A1'] = f'REKAP BULANAN {tahun}'
    ws['A1'].font = Font(bold=True, size=12, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1A3A5C', end_color='1A3A5C', fill_type='solid')
    ws.merge_cells('A1:H1')

    headers = ['No', 'Bulan', 'Meja', 'Pekerjaan', 'Total Target', 'Total Realisasi', 'Progress %', 'Status']
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col, value=h)
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center')

    row_num = 3
    counter = 1
    for bulan in range(1, 13):
        data = get_rekap_bulanan(tahun, bulan, meja)
        for d in data:
            pct = d.get('persentase', 0)
            status = hitung_status(d['total_realisasi'], d['total_target'], threshold_pct)
            ws.cell(row=row_num, column=1, value=counter)
            ws.cell(row=row_num, column=2, value=d.get('nama_bulan', ''))
            ws.cell(row=row_num, column=3, value=d.get('meja', ''))
            ws.cell(row=row_num, column=4, value=d.get('pekerjaan', ''))
            ws.cell(row=row_num, column=5, value=d.get('total_target', 0))
            ws.cell(row=row_num, column=6, value=d.get('total_realisasi', 0))
            c = ws.cell(row=row_num, column=7, value=round(pct / 100, 4))
            c.number_format = '0.00%'
            ws.cell(row=row_num, column=8, value=status)
            row_num += 1
            counter += 1


def _write_stats_sheet(ws, settings):
    """Menulis sheet statistik ringkasan."""
    from app.services.progress_service import get_dashboard_stats, get_perbandingan_meja
    ws['A1'] = f'STATISTIK PROGRESS PLSCS {settings.report_year}'
    ws['A1'].font = Font(bold=True, size=14, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color='1A3A5C', end_color='1A3A5C', fill_type='solid')
    ws.merge_cells('A1:D1')

    stats_all = get_dashboard_stats(tahun=settings.report_year)
    stats_a = get_dashboard_stats(tahun=settings.report_year, meja='Meja A')
    stats_b = get_dashboard_stats(tahun=settings.report_year, meja='Meja B')

    data_rows = [
        ['Parameter', 'Semua', 'Meja A', 'Meja B'],
        ['Total Entri', stats_all.get('total_entri', 0), stats_a.get('total_entri', 0), stats_b.get('total_entri', 0)],
        ['Total Target', stats_all.get('total_target', 0), stats_a.get('total_target', 0), stats_b.get('total_target', 0)],
        ['Total Realisasi', stats_all.get('total_realisasi', 0), stats_a.get('total_realisasi', 0), stats_b.get('total_realisasi', 0)],
        ['Progress %',
         round(stats_all.get('persentase_keseluruhan', 0) / 100, 4),
         round(stats_a.get('persentase_keseluruhan', 0) / 100, 4),
         round(stats_b.get('persentase_keseluruhan', 0) / 100, 4)],
        ['Jumlah Tercapai', stats_all.get('jumlah_tercapai', 0), stats_a.get('jumlah_tercapai', 0), stats_b.get('jumlah_tercapai', 0)],
        ['Jumlah Tidak Tercapai', stats_all.get('jumlah_tidak_tercapai', 0), stats_a.get('jumlah_tidak_tercapai', 0), stats_b.get('jumlah_tidak_tercapai', 0)],
    ]

    for r, row_data in enumerate(data_rows, 2):
        for c, val in enumerate(row_data, 1):
            cell = ws.cell(row=r, column=c, value=val)
            if r == 2:
                cell.fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
                cell.font = Font(bold=True, color='FFFFFF')
            if c > 1 and r == 6:
                cell.number_format = '0.00%'

    ws.column_dimensions['A'].width = 25
    for col in ['B', 'C', 'D']:
        ws.column_dimensions[col].width = 15


def _fmt_date(tanggal: str) -> str:
    """Format tanggal ISO ke format yang lebih mudah dibaca."""
    try:
        d = datetime.strptime(tanggal, '%Y-%m-%d')
        return d.strftime('%d/%m/%Y')
    except Exception:
        return tanggal
